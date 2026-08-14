"""Extract supported source formats into structured evidence records."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import tempfile
from dataclasses import dataclass
from functools import lru_cache
from importlib.metadata import version as package_version
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urljoin, urlparse

import openpyxl
from bs4 import BeautifulSoup, Tag
from llama_index.core import Document
from lxml import etree

from .chunking import build_search_text
from .document_types import (
    chunk_policy_for,
    normalize_document_type,
    resolve_record_kind,
)
from .tokenizer import DEFAULT_MAX_TOKENS, get_embedding_tokenizer


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".xlsm",
    ".csv",
    ".md",
    ".txt",
    ".html",
    ".htm",
    ".xml",
    ".json",
    ".jsonl",
    ".jpg",
    ".jpeg",
    ".png",
    ".tif",
    ".tiff",
}

LEGACY_EXTENSIONS = {".doc", ".xls", ".ppt"}
DETERMINISTIC_EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".json", ".jsonl", ".xml"}
DOCLING_EXTENSIONS = SUPPORTED_EXTENSIONS - DETERMINISTIC_EXTENSIONS
LITERAL_PDF_URL = re.compile(
    r"(?P<url>(?:https?://|/)[^\"'<>]*?\.pdf(?:\?[^\"'<>]*)?)",
    re.IGNORECASE,
)


def file_content_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _base_metadata(
    *,
    document_id: str,
    original_file_name: str,
    stored_file_name: str,
    source_title: str,
    source_url: str,
    source_type: str,
    access_level: str,
    document_version: int,
    content_hash: str,
    file_type: str,
    category: str,
    language: str,
    effective_date: str,
    source_kind: str,
    document_type: str,
    classification_source: str,
    branch_ids: Iterable[str] | None,
) -> dict[str, Any]:
    return {
        "document_id": document_id,
        "kb_document_id": document_id,
        "source_version_id": f"{document_id}:v{document_version}",
        "file_name": original_file_name,
        "original_file_name": original_file_name,
        "stored_file_name": stored_file_name,
        "source_title": source_title or original_file_name,
        "source_url": source_url or "",
        "url": source_url or "",
        "source_type": source_type,
        "access_level": access_level,
        "document_version": document_version,
        "content_hash": content_hash,
        "file_type": file_type,
        "category": category or "",
        "language": language or "",
        "effective_date": effective_date or "",
        "source_kind": source_kind,
        "document_type": document_type,
        "record_kind": resolve_record_kind({"document_type": document_type}),
        "classification_source": classification_source,
        "branch_ids": list(branch_ids or []),
    }


def _make_document(
    evidence_text: str,
    *,
    base_metadata: dict[str, Any],
    section_index: int,
    structural_kind: str,
    locator: dict[str, Any],
    structure_path: Iterable[str] = (),
    extra_metadata: dict[str, Any] | None = None,
) -> Document | None:
    evidence_text = (evidence_text or "").strip()
    if not evidence_text:
        return None

    metadata = {
        **base_metadata,
        "section_index": section_index,
        "structural_kind": structural_kind,
        "structure_path": [part for part in structure_path if part],
        "locator": locator,
        "parser_version": "deterministic-v2",
        **(extra_metadata or {}),
    }
    record_kind = resolve_record_kind(metadata, structural_kind=structural_kind)
    metadata.update({
        "record_kind": record_kind,
        "chunk_policy": chunk_policy_for(record_kind, structural_kind),
        "evidence_text": evidence_text,
    })
    metadata["search_text"] = build_search_text(metadata, evidence_text)

    document = Document(text=evidence_text, metadata=metadata)
    document.id_ = (
        f"{base_metadata['source_version_id']}:record:{section_index}"
    )
    return document


def _field_text(column: Any, value: Any) -> str:
    return f"{str(column).strip()}: {str(value).strip()}"


def _load_csv(path: Path, base_metadata: dict[str, Any]) -> list[Document]:
    documents: list[Document] = []
    with path.open(newline="", encoding="utf-8-sig", errors="strict") as file:
        reader = csv.DictReader(file)
        headers = [str(header) for header in (reader.fieldnames or []) if header is not None]
        header_text = " | ".join(headers)
        for row_index, row in enumerate(reader):
            fields = [
                _field_text(column, value)
                for column, value in row.items()
                if column is not None and value is not None and str(value).strip()
            ]
            document = _make_document(
                "\n".join(fields),
                base_metadata=base_metadata,
                section_index=row_index,
                structural_kind="table_row",
                locator={"type": "table_row", "row": row_index + 2},
                extra_metadata={
                    "row_number": row_index + 2,
                    "table_header": header_text,
                    "repeat_context": header_text,
                },
            )
            if document:
                documents.append(document)
    return documents


def _load_excel(path: Path, base_metadata: dict[str, Any]) -> list[Document]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    documents: list[Document] = []
    section_index = 0
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                raw_headers = next(rows)
            except StopIteration:
                continue
            headers = [str(value).strip() if value is not None else "" for value in raw_headers]
            header_text = " | ".join(header for header in headers if header)
            for row_number, row in enumerate(rows, start=2):
                fields = [
                    _field_text(header, value)
                    for header, value in zip(headers, row)
                    if header and value is not None and str(value).strip()
                ]
                document = _make_document(
                    "\n".join(fields),
                    base_metadata=base_metadata,
                    section_index=section_index,
                    structural_kind="table_row",
                    locator={
                        "type": "sheet_row",
                        "sheet": worksheet.title,
                        "row": row_number,
                    },
                    structure_path=[worksheet.title],
                    extra_metadata={
                        "sheet_name": worksheet.title,
                        "row_number": row_number,
                        "table_header": header_text,
                        "repeat_context": header_text,
                    },
                )
                if document:
                    documents.append(document)
                    section_index += 1
    finally:
        workbook.close()
    return documents


def _json_record_to_text(record: object) -> str:
    if isinstance(record, dict):
        return "\n".join(
            _field_text(
                key,
                json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                if isinstance(value, (dict, list))
                else value,
            )
            for key, value in record.items()
        )
    if isinstance(record, list):
        return "\n".join(str(item) for item in record)
    return str(record)


def _load_json(path: Path, base_metadata: dict[str, Any]) -> list[Document]:
    if path.suffix.lower() == ".jsonl":
        with path.open(encoding="utf-8", errors="strict") as file:
            records = [json.loads(line) for line in file if line.strip()]
    else:
        data = json.loads(path.read_text(encoding="utf-8", errors="strict"))
        records = data if isinstance(data, list) else [data]

    documents: list[Document] = []
    for record_index, record in enumerate(records):
        document = _make_document(
            _json_record_to_text(record),
            base_metadata=base_metadata,
            section_index=record_index,
            structural_kind="record",
            locator={"type": "json_record", "record": record_index},
            extra_metadata={"record_index": record_index},
        )
        if document:
            documents.append(document)
    return documents


def _xml_record_text(element: etree._Element) -> str:
    fields: list[str] = []
    for child in element.iter():
        if child is element and len(element):
            continue
        tag = etree.QName(child.tag).localname
        value = " ".join(part.strip() for part in child.itertext() if part.strip())
        if value:
            fields.append(f"{tag}: {value}")
    if not fields:
        fields.append(" ".join(part.strip() for part in element.itertext() if part.strip()))
    return "\n".join(fields)


def _load_xml(path: Path, base_metadata: dict[str, Any]) -> list[Document]:
    parser = etree.XMLParser(
        recover=False,
        resolve_entities=False,
        no_network=True,
        huge_tree=False,
    )
    root = etree.parse(str(path), parser).getroot()
    targets = list(root) or [root]
    documents: list[Document] = []
    for record_index, element in enumerate(targets):
        tag = etree.QName(element.tag).localname
        document = _make_document(
            _xml_record_text(element),
            base_metadata=base_metadata,
            section_index=record_index,
            structural_kind="record",
            locator={"type": "xml_record", "tag": tag, "record": record_index},
            structure_path=[etree.QName(root.tag).localname, tag],
            extra_metadata={"xml_tag": tag, "record_index": record_index},
        )
        if document:
            documents.append(document)
    return documents


@dataclass(frozen=True)
class FaqPair:
    question: str
    answer: str
    anchor: str = ""


def _clean_html_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, Tag):
        return value.get_text("\n", strip=True)
    if isinstance(value, str) and "<" in value:
        return BeautifulSoup(value, "html.parser").get_text("\n", strip=True)
    return str(value).strip()


def _walk_json(value: Any) -> Iterable[dict[str, Any]]:
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk_json(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_json(child)


def _faq_from_schema_org(soup: BeautifulSoup) -> list[FaqPair]:
    pairs: list[FaqPair] = []
    for script in soup.select('script[type="application/ld+json"]'):
        try:
            data = json.loads(script.string or script.get_text())
        except (json.JSONDecodeError, TypeError):
            continue
        for item in _walk_json(data):
            item_type = item.get("@type")
            types = item_type if isinstance(item_type, list) else [item_type]
            if "FAQPage" not in types:
                continue
            for entity in item.get("mainEntity", []):
                if not isinstance(entity, dict):
                    continue
                question = _clean_html_text(entity.get("name"))
                accepted = entity.get("acceptedAnswer") or {}
                answer = _clean_html_text(
                    accepted.get("text") if isinstance(accepted, dict) else accepted
                )
                if question and answer:
                    pairs.append(FaqPair(question, answer))
    return pairs


def _faq_from_details(soup: BeautifulSoup) -> list[FaqPair]:
    pairs: list[FaqPair] = []
    for details in soup.find_all("details"):
        summary = details.find("summary")
        if not summary:
            continue
        answer_parts = [
            _clean_html_text(child)
            for child in details.children
            if child is not summary and _clean_html_text(child)
        ]
        question = _clean_html_text(summary)
        answer = "\n".join(answer_parts).strip()
        if question and answer:
            pairs.append(FaqPair(question, answer, str(details.get("id") or "")))
    return pairs


def _faq_from_definition_lists(soup: BeautifulSoup) -> list[FaqPair]:
    pairs: list[FaqPair] = []
    for term in soup.find_all("dt"):
        answers: list[str] = []
        sibling = term.find_next_sibling()
        while sibling and sibling.name != "dt":
            if sibling.name == "dd":
                answers.append(_clean_html_text(sibling))
            sibling = sibling.find_next_sibling()
        question = _clean_html_text(term)
        answer = "\n".join(part for part in answers if part)
        if question and answer:
            pairs.append(FaqPair(question, answer, str(term.get("id") or "")))
    return pairs


def _faq_from_aria(soup: BeautifulSoup) -> list[FaqPair]:
    pairs: list[FaqPair] = []
    for control in soup.select("[aria-controls]"):
        target_id = str(control.get("aria-controls") or "").strip()
        target = soup.find(id=target_id) if target_id else None
        question = _clean_html_text(control)
        answer = _clean_html_text(target)
        if question and answer:
            pairs.append(FaqPair(question, answer, target_id))
    return pairs


def _is_question_block(tag: Tag) -> bool:
    text = _clean_html_text(tag)
    if not text.endswith(("?", "？")):
        return False
    if tag.name in {"h1", "h2", "h3", "h4", "h5", "h6", "strong", "b"}:
        return True
    return tag.name in {"p", "div", "li"} and bool(tag.find(["strong", "b", "em"]))


def _faq_from_question_blocks(soup: BeautifulSoup) -> list[FaqPair]:
    pairs: list[FaqPair] = []
    candidates = [tag for tag in soup.find_all(True) if _is_question_block(tag)]
    candidate_ids = {id(tag) for tag in candidates}
    for question_block in candidates:
        sibling = question_block.find_next_sibling()
        answer_parts: list[str] = []
        while sibling:
            if id(sibling) in candidate_ids or sibling.name in {"h1", "h2", "h3"}:
                break
            text = _clean_html_text(sibling)
            if text:
                answer_parts.append(text)
            if answer_parts:
                break
            sibling = sibling.find_next_sibling()
        question = _clean_html_text(question_block)
        answer = "\n".join(answer_parts)
        if question and answer:
            pairs.append(FaqPair(question, answer, str(question_block.get("id") or "")))
    return pairs


_QUESTION_MARKER = re.compile(
    r"^\s*(?:(?:Q(?:uestion)?)(?:\s*\.?\s*\d+)?|問(?:題)?|问题)"
    r"\s*(?:[.:：、)]\s*)?(?P<text>.*)$",
    re.IGNORECASE,
)
_ANSWER_MARKER = re.compile(
    r"^\s*(?:(?:A(?:nswer)?)(?:\s*\.?\s*\d+)?|答(?:案)?)"
    r"\s*(?:[.:：、)]\s*)?(?P<text>.*)$",
    re.IGNORECASE,
)


def _faq_from_markers(soup: BeautifulSoup) -> list[FaqPair]:
    lines = [line.strip() for line in soup.get_text("\n").splitlines() if line.strip()]
    pairs: list[FaqPair] = []
    question = ""
    answer_parts: list[str] = []
    state = ""

    def flush() -> None:
        nonlocal question, answer_parts
        answer = "\n".join(answer_parts).strip()
        if question and answer:
            pairs.append(FaqPair(question, answer))
        question = ""
        answer_parts = []

    for line in lines:
        q_match = _QUESTION_MARKER.match(line)
        a_match = _ANSWER_MARKER.match(line)
        if q_match:
            flush()
            question = q_match.group("text").strip()
            state = "question"
        elif a_match and question:
            answer_parts = [a_match.group("text").strip()] if a_match.group("text").strip() else []
            state = "answer"
        elif state == "question" and question:
            question = f"{question}\n{line}".strip()
        elif state == "answer":
            answer_parts.append(line)
    flush()
    return pairs


def extract_faq_pairs(html: str, *, hinted: bool = False) -> list[FaqPair]:
    """Extract FAQ pairs using ordered structural strategies."""
    soup = BeautifulSoup(html, "html.parser")
    strategies = (
        _faq_from_schema_org,
        _faq_from_details,
        _faq_from_definition_lists,
        _faq_from_aria,
        _faq_from_question_blocks,
        _faq_from_markers,
    )
    for index, strategy in enumerate(strategies):
        pairs = strategy(soup)
        if pairs and (hinted or index < 4 or len(pairs) >= 2):
            return pairs
    return []


def extract_html_record_metadata(html: str, record_kind: str) -> dict[str, Any]:
    """Extract typed date fields without using them as chunk boundaries."""
    soup = BeautifulSoup(html, "html.parser")
    metadata: dict[str, Any] = {}
    wanted_types = {
        "record": {"Event", "NewsArticle", "Article"},
        "event": {"Event"},
        "notice": {"NewsArticle", "Article"},
    }
    date_fields = {
        "record": ("startDate", "endDate", "datePublished", "dateModified"),
        "event": ("startDate", "endDate"),
        "notice": ("datePublished", "dateModified"),
    }
    for script in soup.select('script[type="application/ld+json"]'):
        try:
            data = json.loads(script.string or script.get_text())
        except (json.JSONDecodeError, TypeError):
            continue
        for item in _walk_json(data):
            item_type = item.get("@type")
            item_types = set(item_type if isinstance(item_type, list) else [item_type])
            if not item_types.intersection(wanted_types.get(record_kind, set())):
                continue
            for field in date_fields.get(record_kind, ()):
                if item.get(field):
                    metadata[field] = str(item[field])
            if item.get("name"):
                metadata["record_header"] = _clean_html_text(item["name"])

    time_values = [
        str(tag.get("datetime"))
        for tag in soup.find_all("time")
        if tag.get("datetime")
    ]
    if time_values:
        metadata["html_time_values"] = time_values
    if not metadata.get("effective_date"):
        metadata["effective_date"] = (
            metadata.get("startDate")
            or metadata.get("datePublished")
            or (time_values[0] if time_values else "")
        )
    return metadata


def _load_html_faq(
    path: Path,
    base_metadata: dict[str, Any],
) -> list[Document]:
    html = path.read_text(encoding="utf-8", errors="strict")
    hinted = base_metadata.get("record_kind") == "faq"
    pairs = extract_faq_pairs(html, hinted=hinted)
    if not pairs:
        return []

    documents: list[Document] = []
    parent_record = hashlib.sha256(
        str(base_metadata["source_version_id"]).encode("utf-8")
    ).hexdigest()[:24]
    for pair_index, pair in enumerate(pairs):
        evidence = f"{pair.question}\n{pair.answer}"
        locator = {
            "type": "web_anchor" if pair.anchor else "faq_record",
            "anchor": pair.anchor,
            "record": pair_index,
        }
        document = _make_document(
            evidence,
            base_metadata={**base_metadata, "record_kind": "faq"},
            section_index=pair_index,
            structural_kind="faq_pair",
            locator=locator,
            extra_metadata={
                "question": pair.question,
                "record_header": pair.question,
                "repeat_context": pair.question,
                "parent_record_id": f"{parent_record}:{pair_index}",
                "answer_text": pair.answer,
            },
        )
        if document:
            documents.append(document)
    return documents


@lru_cache(maxsize=4)
def _docling_converter(ocr_languages: str):
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.object_detection_engine_options import (
        TransformersObjectDetectionEngineOptions,
    )
    from docling.datamodel.pipeline_options import (
        LayoutObjectDetectionOptions,
        PdfPipelineOptions,
        TesseractCliOcrOptions,
    )
    from docling.document_converter import DocumentConverter, PdfFormatOption

    artifacts_path = Path(os.getenv("DOCLING_ARTIFACTS_PATH", "/app/models/docling"))
    if not artifacts_path.exists():
        raise RuntimeError(
            f"Docling artifacts are not available at {artifacts_path}; ingestion cannot continue offline."
        )

    pipeline_options = PdfPipelineOptions(
        artifacts_path=artifacts_path,
        enable_remote_services=False,
        allow_external_plugins=False,
        document_timeout=float(os.getenv("DOCLING_DOCUMENT_TIMEOUT_SECONDS", "300")),
        do_ocr=True,
        ocr_options=TesseractCliOcrOptions(
            lang=[language for language in ocr_languages.split("+") if language]
        ),
        do_table_structure=True,
        do_code_enrichment=False,
        do_formula_enrichment=False,
        do_picture_classification=False,
        do_picture_description=False,
        # torch.compile needs a C++ toolchain at runtime. Compilation is only a
        # performance optimization and makes the slim offline image brittle.
        layout_options=LayoutObjectDetectionOptions.from_preset(
            "layout_heron_default",
            engine_options=TransformersObjectDetectionEngineOptions(
                compile_model=False,
            ),
        ),
    )
    format_option = PdfFormatOption(pipeline_options=pipeline_options)
    return DocumentConverter(format_options={
        InputFormat.PDF: format_option,
        InputFormat.IMAGE: format_option,
    })


def _docling_json_path(base_metadata: dict[str, Any]) -> Path:
    output_root = Path(os.getenv("DOCLING_JSON_DIR", "/app/storage/docling"))
    output_root.mkdir(parents=True, exist_ok=True)
    return output_root / (
        f"{base_metadata['document_id']}-v{base_metadata['document_version']}-"
        f"{base_metadata['content_hash'][:12]}.json"
    )


def _persist_docling_document(
    docling_document: Any,
    base_metadata: dict[str, Any],
) -> Path:
    output_path = _docling_json_path(base_metadata)
    payload = {
        "parser_version": f"docling-{package_version('docling')}",
        "source_version_id": base_metadata["source_version_id"],
        "content_hash": base_metadata["content_hash"],
        "document": docling_document.model_dump(mode="json"),
    }
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    return output_path


def _load_or_convert_docling(
    path: Path,
    base_metadata: dict[str, Any],
    ocr_languages: str,
) -> tuple[Any, Path]:
    from docling_core.types.doc import DoclingDocument

    json_path = _docling_json_path(base_metadata)
    parser_version = f"docling-{package_version('docling')}"
    if json_path.exists():
        payload = json.loads(json_path.read_text(encoding="utf-8"))
        if (
            payload.get("content_hash") == base_metadata["content_hash"]
            and payload.get("parser_version") == parser_version
        ):
            return DoclingDocument.model_validate(payload["document"]), json_path

    conversion_path = path
    temporary_directory: tempfile.TemporaryDirectory[str] | None = None
    if path.suffix.lower() in {".html", ".htm"}:
        html = path.read_text(encoding="utf-8", errors="strict")
        if not re.search(r"<html(?:\s|>)", html, re.IGNORECASE):
            temporary_directory = tempfile.TemporaryDirectory()
            conversion_path = Path(temporary_directory.name) / path.name
            conversion_path.write_text(
                f"<!doctype html><html><body>{html}</body></html>",
                encoding="utf-8",
            )
    try:
        result = _docling_converter(ocr_languages).convert(conversion_path)
    finally:
        if temporary_directory is not None:
            temporary_directory.cleanup()
    docling_document = result.document
    return docling_document, _persist_docling_document(docling_document, base_metadata)


def _docling_locator(doc_items: list[Any], structure_path: list[str]) -> dict[str, Any]:
    references = [str(item.self_ref) for item in doc_items if getattr(item, "self_ref", None)]
    pages = sorted({
        int(prov.page_no)
        for item in doc_items
        for prov in (getattr(item, "prov", None) or [])
        if getattr(prov, "page_no", None) is not None
    })
    locator: dict[str, Any] = {
        "type": "docling_items",
        "item_refs": references,
    }
    if pages:
        locator.update({
            "type": "page" if len(pages) == 1 else "page_range",
            "page": pages[0],
            "page_end": pages[-1],
        })
    if structure_path:
        locator["heading_path"] = structure_path
    return locator


def _table_header_text(table_text: str) -> str:
    lines = [line for line in table_text.splitlines() if line.strip()]
    if not lines:
        return ""
    if len(lines) > 1 and re.fullmatch(r"[\s|:+-]+", lines[1]):
        return "\n".join(lines[:2])
    return lines[0]


def _load_html_form_directory(
    path: Path,
    base_metadata: dict[str, Any],
) -> list[Document]:
    """Turn the HKPL forms landing page into searchable form-link records."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="strict"), "html.parser")
    documents: list[Document] = []
    for row_index, row in enumerate(soup.select("table tr")):
        cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
        link = row.find("a", href=True)
        if len(cells) < 2 or not link or cells[:2] == ["Form No.", "Form Name"]:
            continue
        form_number, form_name = cells[:2]
        href = urljoin(
            str(base_metadata.get("source_url") or ""),
            str(link.get("href") or "").strip(),
        )
        evidence = f"Form number: {form_number}\nForm name: {form_name}\nDownload URL: {href}"
        document = _make_document(
            evidence,
            base_metadata=base_metadata,
            section_index=len(documents),
            structural_kind="table_row",
            locator={"type": "web_link", "row": row_index + 1, "url": href},
            structure_path=("Forms",),
            extra_metadata={
                "form_number": form_number,
                "form_name": form_name,
                "download_url": href,
                "table_header": "Form number | Form name | Download URL",
                "repeat_context": "Form number | Form name | Download URL",
            },
        )
        if document:
            documents.append(document)
    return documents


def _load_html_fallback(
    path: Path,
    base_metadata: dict[str, Any],
) -> list[Document]:
    """Recover useful HTML text when Docling exports an empty document."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="strict"), "html.parser")
    for tag in soup.select("script, style, noscript, template"):
        tag.decompose()

    base_url = str(base_metadata.get("source_url") or "")
    base_host = urlparse(base_url).netloc.lower()
    resource_links: list[tuple[str, str]] = []
    seen_urls: set[str] = set()
    for tag in soup.find_all(True):
        label = re.sub(r"\s+", " ", tag.get_text(" ", strip=True)).strip()
        for value in tag.attrs.values():
            values = value if isinstance(value, list) else [value]
            for item in values:
                for match in LITERAL_PDF_URL.finditer(str(item)):
                    absolute = urljoin(base_url, match.group("url"))
                    parsed = urlparse(absolute)
                    if (
                        parsed.scheme not in {"http", "https"}
                        or (base_host and parsed.netloc.lower() != base_host)
                        or absolute in seen_urls
                    ):
                        continue
                    seen_urls.add(absolute)
                    resource_links.append((label or Path(parsed.path).name, absolute))

    lines = list(dict.fromkeys(
        re.sub(r"\s+", " ", line).strip()
        for line in soup.get_text("\n").splitlines()
        if line.strip()
    ))
    lines.extend(f"{label}: {url}" for label, url in resource_links)
    evidence = "\n".join(lines).strip()

    # Empty Docling output plus interactive controls is normally a login/search
    # shell. A genuine resource directory remains eligible through its links.
    if (soup.select_one("form, input, textarea, select") and not resource_links) or (
        len(evidence) < 200 and not resource_links
    ):
        return []

    document = _make_document(
        evidence,
        base_metadata=base_metadata,
        section_index=0,
        structural_kind="fallback",
        locator={"type": "web_document", "url": base_url},
        extra_metadata={"parser_version": "html-fallback-v1"},
    )
    return [document] if document else []


def _load_html_event_records(
    path: Path,
    base_metadata: dict[str, Any],
) -> list[Document]:
    """Keep each HKPL event-detail table as one field-aware record."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="strict"), "html.parser")
    root = soup.select_one(".main_content") or soup
    documents: list[Document] = []
    for table_index, table in enumerate(root.find_all("table")):
        fields: list[str] = []
        for row in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])]
            if len(cells) >= 2 and any(cells):
                fields.append(f"{cells[0].rstrip(':')}: {' | '.join(cells[1:])}")
        evidence = "\n".join(fields)
        document = _make_document(
            evidence,
            base_metadata=base_metadata,
            section_index=len(documents),
            structural_kind="record",
            locator={"type": "web_table", "table": table_index + 1},
            structure_path=(base_metadata.get("source_title") or "Event",),
        )
        if document:
            documents.append(document)
    return documents


def _load_html_branch_profile(
    path: Path,
    base_metadata: dict[str, Any],
) -> list[Document]:
    """Preserve each tab of an HKPL branch profile as an addressable section."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="strict"), "html.parser")
    documents: list[Document] = []
    for panel in soup.select(".info_tabcont"):
        evidence = panel.get_text("\n", strip=True)
        if len(evidence) < 40 or "loading events" in evidence.lower():
            continue
        anchor = str(panel.get("id") or "")
        document = _make_document(
            evidence,
            base_metadata=base_metadata,
            section_index=len(documents),
            structural_kind="record",
            locator={"type": "web_anchor", "anchor": anchor},
            structure_path=(base_metadata.get("source_title") or "Library profile",),
        )
        if document:
            documents.append(document)
    return documents


def _expanded_table_rows(table: Tag) -> list[list[str]]:
    """Expand row/column spans so schedule cells retain their row context."""
    rows: list[list[str]] = []
    pending: dict[int, tuple[str, int]] = {}
    for html_row in table.find_all("tr"):
        values: dict[int, str] = {}
        for column, (value, remaining) in list(pending.items()):
            values[column] = value
            if remaining <= 1:
                del pending[column]
            else:
                pending[column] = (value, remaining - 1)

        column = 0
        for cell in html_row.find_all(["th", "td"], recursive=False):
            while column in values:
                column += 1
            value = cell.get_text(" ", strip=True)
            colspan = max(int(cell.get("colspan") or 1), 1)
            rowspan = max(int(cell.get("rowspan") or 1), 1)
            for offset in range(colspan):
                values[column + offset] = value
                if rowspan > 1:
                    pending[column + offset] = (value, rowspan - 1)
            column += colspan
        if values:
            rows.append([values.get(index, "") for index in range(max(values) + 1)])
    return rows


def _load_html_hours_rows(
    path: Path,
    base_metadata: dict[str, Any],
) -> list[Document]:
    """Index HKPL hours and mobile schedules by logical table row."""
    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="strict"), "html.parser")
    root = soup.select_one(".main_content") or soup
    documents: list[Document] = []
    for table_index, table in enumerate(root.find_all("table")):
        rows = _expanded_table_rows(table)
        if not rows:
            continue
        wide = max(len(row) for row in rows) >= 3
        header_count = 0
        if wide:
            header_count = 2 if len(rows) > 1 and any(
                "week " in cell.lower() for cell in rows[1]
            ) else 1
        header = "\n".join(" | ".join(row) for row in rows[:header_count]).strip()
        group = ""
        for row_index, row in enumerate(rows[header_count:], start=header_count + 1):
            unique_values = [value for value in dict.fromkeys(row) if value]
            if wide and len(unique_values) == 1:
                group = unique_values[0]
                continue
            row_text = " | ".join(row).strip(" |")
            evidence = "\n".join(part for part in (header, group, row_text) if part)
            document = _make_document(
                evidence,
                base_metadata=base_metadata,
                section_index=len(documents),
                structural_kind="table_row",
                locator={"type": "web_table_row", "table": table_index + 1, "row": row_index},
                structure_path=(base_metadata.get("source_title") or "Opening hours",),
                extra_metadata={
                    "table_header": header,
                    "repeat_context": header,
                },
            )
            if document:
                documents.append(document)
    return documents


def _load_docling(
    path: Path,
    base_metadata: dict[str, Any],
    ocr_languages: str,
) -> list[Document]:
    from docling_core.transforms.chunker.hybrid_chunker import HybridChunker
    from docling_core.types.doc import TableItem

    docling_document, json_path = _load_or_convert_docling(
        path,
        base_metadata,
        ocr_languages,
    )
    parser_version = f"docling-{package_version('docling')}"
    source_parent_id = hashlib.sha256(
        str(base_metadata["source_version_id"]).encode("utf-8")
    ).hexdigest()[:24]

    if normalize_document_type(base_metadata.get("document_type")) == "auto":
        evidence_text = docling_document.export_to_markdown().strip()
        if not evidence_text:
            return []
        metadata = {
            **base_metadata,
            "section_index": 0,
            "structural_kind": "fallback",
            "structure_path": [],
            "locator": {"type": "document"},
            "parser_version": parser_version,
            "docling_json_path": str(json_path),
            "parent_record_id": source_parent_id,
            "chunk_policy": "fallback",
            "evidence_text": evidence_text,
        }
        metadata["search_text"] = build_search_text(metadata, evidence_text)
        return [Document(text=evidence_text, metadata=metadata)]

    tokenizer = get_embedding_tokenizer(DEFAULT_MAX_TOKENS)
    chunker = HybridChunker(
        tokenizer=tokenizer,
        repeat_table_header=True,
        merge_peers=True,
        omit_header_on_overflow=False,
    )
    documents: list[Document] = []
    for chunk_index, chunk in enumerate(chunker.chunk(dl_doc=docling_document)):
        evidence_text = str(chunk.text or "").strip()
        if not evidence_text:
            continue
        structure_path = [str(value) for value in (chunk.meta.headings or []) if value]
        doc_items = list(chunk.meta.doc_items or [])
        is_table = any(isinstance(item, TableItem) for item in doc_items)
        structural_kind = "table" if is_table else "hierarchical_leaf"
        locator = _docling_locator(doc_items, structure_path)
        metadata = {
            **base_metadata,
            "section_index": chunk_index,
            "structural_kind": structural_kind,
            "structure_path": structure_path,
            "locator": locator,
            "parser_version": parser_version,
            "docling_json_path": str(json_path),
            "docling_item_refs": locator.get("item_refs", []),
            "parent_record_id": source_parent_id,
        }
        record_kind = resolve_record_kind(metadata, structural_kind=structural_kind)
        if record_kind == "record":
            record_header = str(
                base_metadata.get("record_header")
                or base_metadata.get("source_title")
                or ""
            ).strip()
            metadata.update({
                "record_header": record_header,
                "repeat_context": record_header,
            })
            if record_header and not evidence_text.startswith(record_header):
                evidence_text = f"{record_header}\n{evidence_text}"
        if is_table:
            table_header = _table_header_text(evidence_text)
            metadata.update({
                "table_header": table_header,
                "repeat_context": table_header,
            })
        metadata.update({
            "record_kind": record_kind,
            "chunk_policy": chunk_policy_for(
                str(base_metadata.get("document_type") or "auto"),
                structural_kind,
            ),
            "evidence_text": evidence_text,
        })
        if locator.get("page"):
            metadata["page_number"] = locator["page"]
        if structure_path:
            metadata["section_heading"] = structure_path[-1]
        metadata["search_text"] = build_search_text(metadata, evidence_text)

        document = Document(text=evidence_text, metadata=metadata)
        document.id_ = f"{base_metadata['source_version_id']}:docling:{chunk_index}"
        documents.append(document)
    return documents


def load_file(
    path: Path,
    *,
    document_id: str,
    original_file_name: str,
    source_title: str = "",
    source_url: str = "",
    source_type: str = "admin_upload",
    access_level: str = "public",
    document_version: int = 1,
    content_hash: str | None = None,
    ocr_languages: str = "eng+chi_tra+chi_sim",
    category: str | None = None,
    language: str | None = None,
    effective_date: str | None = None,
    source_kind: str = "upload",
    document_type: str = "auto",
    classification_source: str | None = None,
    branch_ids: Iterable[str] | None = None,
) -> list[Document]:
    """Dispatch a source to a deterministic reader or pinned Docling parser."""
    extension = path.suffix.lower()
    if extension in LEGACY_EXTENSIONS:
        raise ValueError(
            f"Legacy format {extension} is not supported directly. "
            "Convert it to .docx, .xlsx, or .pptx first."
        )
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file extension: {extension}")

    actual_hash = content_hash or file_content_hash(path)
    selected_type = normalize_document_type(document_type)
    base_metadata = _base_metadata(
        document_id=document_id,
        original_file_name=original_file_name,
        stored_file_name=path.name,
        source_title=source_title,
        source_url=source_url,
        source_type=source_type,
        access_level=access_level,
        document_version=document_version,
        content_hash=actual_hash,
        file_type=extension.lstrip("."),
        category=category or "",
        language=language or "",
        effective_date=effective_date or "",
        source_kind=source_kind,
        document_type=document_type,
        classification_source=(
            classification_source
            or ("fallback" if selected_type == "auto" else "librarian")
        ),
        branch_ids=branch_ids,
    )

    if extension == ".csv":
        return _load_csv(path, base_metadata)
    if extension in {".xlsx", ".xlsm"}:
        return _load_excel(path, base_metadata)
    if extension in {".json", ".jsonl"}:
        return _load_json(path, base_metadata)
    if extension == ".xml":
        return _load_xml(path, base_metadata)
    if extension in {".html", ".htm"}:
        html = path.read_text(encoding="utf-8", errors="strict")
        html_source_url = str(base_metadata.get("source_url") or "")
        if re.search(
            r"/about-us/forms\.html(?:[?#].*)?$",
            html_source_url,
            re.IGNORECASE,
        ):
            form_documents = _load_html_form_directory(path, base_metadata)
            if form_documents:
                return form_documents
        if re.search(
            r"/extension-activities/(?:event|sub-event)/\d+",
            html_source_url,
            re.IGNORECASE,
        ):
            event_documents = _load_html_event_records(path, base_metadata)
            if event_documents:
                return event_documents
        if re.search(
            r"/locations/opening-hours(?:-\d+)?\.html(?:[?#].*)?$",
            html_source_url,
            re.IGNORECASE,
        ):
            hours_documents = _load_html_hours_rows(path, base_metadata)
            if hours_documents:
                return hours_documents
        if re.search(
            r"/locations/(?!opening-hours|mobile-libraries|libraries)[^/]+/[^/]+\.html(?:[?#].*)?$",
            html_source_url,
            re.IGNORECASE,
        ):
            branch_documents = _load_html_branch_profile(path, base_metadata)
            if branch_documents:
                return branch_documents
        if base_metadata["record_kind"] == "record":
            base_metadata.update(
                extract_html_record_metadata(html, base_metadata["record_kind"])
            )
        if base_metadata["record_kind"] == "faq":
            faq_documents = _load_html_faq(path, base_metadata)
            if faq_documents:
                return faq_documents
    if extension in DOCLING_EXTENSIONS:
        documents = _load_docling(path, base_metadata, ocr_languages)
        if documents or extension not in {".html", ".htm"}:
            return documents
        return _load_html_fallback(path, base_metadata)
    raise ValueError(f"No reader configured for {extension}")
