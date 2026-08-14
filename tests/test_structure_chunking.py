from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from scripts.crawl_hkpl_site import (
    decode_response_text,
    discover_links,
    extract_main_html,
    flush_pending,
)

import openpyxl
from llama_index.core import Document

import src.ingestion.service as ingestion_service
from src.ingestion.chunking import build_search_text, chunk_documents
from src.ingestion.classification import (
    MAX_BATCH_ITEMS,
    classification_sample,
    classify_batch_items,
    classify_batch_items_resilient,
)
from src.ingestion.document_types import (
    chunk_policy_for,
    resolve_record_kind,
    validate_document_type,
)
from src.ingestion.readers import (
    extract_faq_pairs,
    extract_html_record_metadata,
    load_file,
)


FIXTURES = Path(__file__).parent / "fixtures"


class CharacterTokenizer:
    def __init__(self, max_tokens: int = 512):
        self.max_tokens = max_tokens

    def count_tokens(self, text: str) -> int:
        return len(text)

    def get_max_tokens(self) -> int:
        return self.max_tokens

    def get_tokenizer(self):
        return self

    def encode(self, text: str, **_: object) -> list[int]:
        return list(range(len(text)))

    def __call__(self, text: str, **_: object) -> dict:
        return {"offset_mapping": [(index, index + 1) for index in range(len(text))]}


def source_document(
    evidence_text: str,
    *,
    locator: dict | None = None,
    structural_kind: str = "hierarchical_leaf",
    record_kind: str = "prose",
    **metadata: object,
) -> Document:
    base = {
        "document_id": "doc-1",
        "kb_document_id": "doc-1",
        "source_version_id": "doc-1:v1",
        "document_version": 1,
        "source_title": "HKPL Source",
        "document_type": record_kind,
        "record_kind": record_kind,
        "structural_kind": structural_kind,
        "chunk_policy": chunk_policy_for(record_kind, structural_kind),
        "structure_path": [],
        "locator": locator or {"type": "section", "section": 1},
        "parser_version": "test-parser",
        "evidence_text": evidence_text,
        **metadata,
    }
    return Document(text=evidence_text, metadata=base)


class DocumentTypeTests(unittest.TestCase):
    def test_physical_table_precedes_admin_faq_hint(self) -> None:
        metadata = {"document_type": "faq", "structural_kind": "table_row"}
        self.assertEqual(resolve_record_kind(metadata), "table")
        self.assertEqual(chunk_policy_for("faq", "table_row"), "table_rows")

    def test_unlabelled_input_does_not_guess_from_content(self) -> None:
        self.assertEqual(resolve_record_kind({"document_type": "auto"}), "prose")

    def test_legacy_admin_values_remain_valid_hints(self) -> None:
        self.assertEqual(validate_document_type("announcement"), "announcement")
        self.assertEqual(resolve_record_kind({"document_type": "announcement"}), "record")
        self.assertEqual(resolve_record_kind({"document_type": "directory"}), "record")


class IngestionHandoffTests(unittest.TestCase):
    def test_registered_document_uses_configured_embedding_model(self) -> None:
        record = {
            "document_id": "doc-1",
            "original_file_name": "source.html",
            "document_type": "prose",
            "version": 1,
        }
        documents = [source_document("Official library information.")]
        nodes = [object()]
        storage_context = object()
        with (
            patch.object(ingestion_service, "ensure_corpus_writable"),
            patch.object(ingestion_service, "ensure_registry_schema"),
            patch.object(ingestion_service, "get_document", return_value=record),
            patch.object(ingestion_service, "_load_record", return_value=documents),
            patch.object(ingestion_service, "chunk_documents", return_value=nodes),
            patch.object(
                ingestion_service.StorageContext,
                "from_defaults",
                return_value=storage_context,
            ),
            patch.object(ingestion_service, "VectorStoreIndex") as vector_index,
            patch.object(ingestion_service, "delete_old_versions", return_value=0),
            patch.object(ingestion_service, "update_status"),
        ):
            result = ingestion_service.process_registered_document("doc-1")

        self.assertEqual(result["status"], "completed")
        self.assertEqual(vector_index.call_args.args[0], nodes)
        self.assertIs(
            vector_index.call_args.kwargs["embed_model"],
            ingestion_service.embed_model,
        )
        self.assertIs(
            vector_index.call_args.kwargs["storage_context"],
            storage_context,
        )


class BatchClassificationTests(unittest.IsolatedAsyncioTestCase):
    async def test_uses_one_compact_generation_call(self) -> None:
        calls: list[tuple[str, dict]] = []

        async def fake_llm(prompt: str, **kwargs: object) -> str:
            calls.append((prompt, kwargs))
            return json.dumps({"items": [
                {"id": "0", "type": "faq"},
                {"id": "1", "type": "skip"},
            ]})

        result = await classify_batch_items([
            {"id": "a", "title": "FAQ", "text": "Q: Renew? A: Online."},
            {"id": "b", "title": "Index", "text": "Links to notices"},
        ], llm_call=fake_llm)

        self.assertEqual(result, {
            "a": {"document_type": "faq"},
            "b": {"document_type": "skip"},
        })
        self.assertEqual(len(calls), 1)
        prompt, options = calls[0]
        self.assertIn("faq|record|prose|skip", prompt)
        self.assertIn("factual tables, directories", prompt)
        self.assertIn("Keep service/guidance pages", prompt)
        self.assertEqual(options["temperature"], 0.0)
        self.assertFalse(options["enable_thinking"])
        self.assertEqual(options["response_format"]["type"], "json_schema")
        self.assertTrue(options["response_format"]["json_schema"]["strict"])
        id_schema = options["response_format"]["json_schema"]["schema"][
            "properties"
        ]["items"]["items"]["properties"]["id"]
        self.assertEqual(id_schema["enum"], ["0", "1"])

    async def test_rejects_partial_or_invalid_output(self) -> None:
        items = [{"id": "a", "text": "one"}, {"id": "b", "text": "two"}]
        invalid_outputs = (
            '{"items":[{"id":"0","type":"prose"}]}',
            '{"items":[{"id":"0","type":"prose"},{"id":"x","type":"faq"}]}',
            '{"items":[{"id":"0","type":"prose"},{"id":"0","type":"faq"}]}',
            '{"items":[{"id":"0","type":"event"},{"id":"1","type":"prose"}]}',
        )
        for raw in invalid_outputs:
            async def fake_llm(*_: object, output: str = raw, **__: object) -> str:
                return output

            with self.subTest(raw=raw), self.assertRaises(ValueError):
                await classify_batch_items(items, llm_call=fake_llm)

    async def test_rejects_oversized_batch_without_a_model_call(self) -> None:
        called = False

        async def fake_llm(*_: object, **__: object) -> str:
            nonlocal called
            called = True
            return "{}"

        with self.assertRaises(ValueError):
            await classify_batch_items(
                [{"id": str(index)} for index in range(MAX_BATCH_ITEMS + 1)],
                llm_call=fake_llm,
            )
        self.assertFalse(called)

    async def test_failed_batch_isolated_to_rejected_item(self) -> None:
        calls = 0

        async def fake_llm(prompt: str, **_: object) -> str:
            nonlocal calls
            calls += 1
            if '"text":"one"' in prompt and '"text":"two"' in prompt:
                raise ValueError("bad batch JSON")
            if '"text":"two"' in prompt:
                raise ValueError("bad item JSON")
            return '{"items":[{"id":"0","type":"prose"}]}'

        decisions = await classify_batch_items_resilient(
            [{"id": "a", "text": "one"}, {"id": "b", "text": "two"}],
            llm_call=fake_llm,
        )

        self.assertEqual(decisions["a"], {"document_type": "prose"})
        self.assertEqual(decisions["b"]["document_type"], "prose")
        self.assertEqual(decisions["b"]["classification_source"], "fallback")
        self.assertEqual(decisions["b"]["classification_error"], "bad item JSON")
        self.assertEqual(calls, 3)

    async def test_llm_decision_precedes_faq_fallback(self) -> None:
        async def fake_llm(*_: object, **__: object) -> str:
            return '{"items":[{"id":"0","type":"prose"}]}'

        decisions = await classify_batch_items_resilient([{
            "id": "faq-page",
            "text": "Q: First?\nA: One\nQ: Second?\nA: Two",
        }], llm_call=fake_llm)

        self.assertEqual(decisions, {"faq-page": {"document_type": "prose"}})

    def test_classifier_sample_keeps_head_and_tail(self) -> None:
        sample = classification_sample("HEAD" + ("x" * 2000) + "TAIL", limit=100)
        self.assertTrue(sample.startswith("HEAD"))
        self.assertTrue(sample.endswith("TAIL"))
        self.assertIn("middle omitted", sample)


class FaqAdapterTests(unittest.TestCase):
    def test_schema_org_has_highest_priority(self) -> None:
        html = """
        <script type="application/ld+json">{
          "@type":"FAQPage","mainEntity":[{
            "@type":"Question","name":"Schema question?",
            "acceptedAnswer":{"@type":"Answer","text":"Schema answer"}
          }]
        }</script>
        <details><summary>Details question?</summary><p>Details answer</p></details>
        """
        pairs = extract_faq_pairs(html)
        self.assertEqual([(pair.question, pair.answer) for pair in pairs], [
            ("Schema question?", "Schema answer")
        ])

    def test_details_multilingual_fixture(self) -> None:
        pairs = extract_faq_pairs(
            (FIXTURES / "hkpl_multilingual_faq.html").read_text(encoding="utf-8")
        )
        self.assertEqual(len(pairs), 2)
        self.assertIn("點樣", pairs[0].question)
        self.assertIn("電子書", pairs[1].question)

    def test_definition_list(self) -> None:
        pairs = extract_faq_pairs("<dl><dt>Borrowing?</dt><dd>Use a library card.</dd></dl>")
        self.assertEqual(pairs[0].answer, "Use a library card.")

    def test_aria_accordion(self) -> None:
        html = '<button aria-controls="answer-1">Opening hours?</button><div id="answer-1">See the branch page.</div>'
        pairs = extract_faq_pairs(html)
        self.assertEqual((pairs[0].question, pairs[0].answer), (
            "Opening hours?", "See the branch page."
        ))

    def test_heading_like_questions_require_repeated_evidence(self) -> None:
        html = """
        <h3>Can I renew?</h3><p>Yes, subject to the lending rules.</p>
        <h3>可以續借嗎？</h3><p>可以，但須符合借閱規則。</p>
        """
        self.assertEqual(len(extract_faq_pairs(html)), 2)
        self.assertEqual(
            extract_faq_pairs("<h3>Rhetorical question?</h3><p>Article prose.</p>"),
            [],
        )

    def test_all_marker_variants(self) -> None:
        html = """<pre>
Q: First?\nA: One
Q. Second?\nAnswer: Two
Q1: Third?\nA1: Three
Q.1: Fourth?\nA.1: Four
問題：第五？\n答案：五
问题：第六？\n答案：六
問：第七？\n答：七
</pre>"""
        pairs = extract_faq_pairs(html, hinted=True)
        self.assertEqual(len(pairs), 7)
        self.assertEqual(pairs[-1].answer, "七")


class ChunkConstructionTests(unittest.TestCase):
    def test_exact_evidence_and_normalized_search_text_are_separate(self) -> None:
        evidence = "ＡＢＣ\u00a0資料"
        document = source_document(
            evidence,
            operational_hash="must-not-be-embedded",
            access_level="public",
        )
        node = chunk_documents(
            [document], tokenizer=CharacterTokenizer(), max_tokens=512
        )[0]
        self.assertEqual(node.metadata["evidence_text"], evidence)
        self.assertIn("ABC 資料", node.text)
        self.assertNotIn("must-not-be-embedded", node.text)
        self.assertNotIn("public", node.text)

    def test_short_meaningful_record_is_retained(self) -> None:
        nodes = chunk_documents(
            [source_document("閉館")], tokenizer=CharacterTokenizer(), max_tokens=512
        )
        self.assertEqual(len(nodes), 1)

    def test_equal_text_at_distinct_locators_is_not_deduplicated(self) -> None:
        documents = [
            source_document("Same official wording", locator={"type": "page", "page": 1}),
            source_document("Same official wording", locator={"type": "page", "page": 2}),
        ]
        nodes = chunk_documents(documents, tokenizer=CharacterTokenizer(), max_tokens=512)
        self.assertEqual(len(nodes), 2)
        self.assertNotEqual(nodes[0].node_id, nodes[1].node_id)

    def test_identical_stable_chunk_ids_are_deduplicated(self) -> None:
        documents = [
            source_document("Same official wording"),
            source_document("Same official wording"),
        ]

        nodes = chunk_documents(
            documents,
            tokenizer=CharacterTokenizer(),
            max_tokens=512,
        )

        self.assertEqual(len(nodes), 1)

    def test_equal_text_and_locator_in_distinct_sections_get_distinct_ids(self) -> None:
        documents = [
            source_document("Same wording", section_index=1),
            source_document("Same wording", section_index=2),
        ]
        nodes = chunk_documents(documents, tokenizer=CharacterTokenizer(), max_tokens=512)
        self.assertNotEqual(nodes[0].node_id, nodes[1].node_id)
        self.assertNotEqual(
            nodes[0].metadata["parent_record_id"],
            nodes[1].metadata["parent_record_id"],
        )

    def test_oversized_prose_prefers_semantic_boundaries_without_overlap(self) -> None:
        document = source_document(
            "First paragraph ends here.\n\n"
            "Second paragraph also ends cleanly.\n\n"
            "Third paragraph finishes naturally.",
            source_title="T",
        )
        nodes = chunk_documents(
            [document], tokenizer=CharacterTokenizer(70), max_tokens=70
        )
        self.assertGreater(len(nodes), 1)
        self.assertTrue(all(node.metadata["chunk_overlap"] == 0 for node in nodes))
        self.assertTrue(all(
            node.metadata["evidence_text"].endswith(".") for node in nodes
        ))

    def test_oversized_faq_repeats_question_and_respects_cap(self) -> None:
        question = "How do I renew a book?"
        document = source_document(
            f"{question}\n" + ("answer " * 80),
            record_kind="faq",
            structural_kind="faq_pair",
            question=question,
            answer_text="answer " * 80,
            record_header=question,
            repeat_context=question,
        )
        nodes = chunk_documents(
            [document], tokenizer=CharacterTokenizer(160), max_tokens=160
        )
        self.assertGreater(len(nodes), 1)
        for node in nodes:
            self.assertLessEqual(node.metadata["token_count"], 160)
            self.assertTrue(node.metadata["evidence_text"].startswith(question))
            self.assertEqual(node.metadata["chunk_policy"], "oversized_leaf")
            self.assertEqual(node.metadata["chunk_overlap"], 64)

    def test_oversized_table_repeats_header(self) -> None:
        header = "Branch | Address"
        document = source_document(
            f"{header}\n" + ("Central Library | Causeway Bay " * 20),
            record_kind="table",
            structural_kind="table_row",
            table_header=header,
            repeat_context=header,
        )
        nodes = chunk_documents(
            [document], tokenizer=CharacterTokenizer(160), max_tokens=160
        )
        self.assertGreater(len(nodes), 1)
        self.assertTrue(all(node.metadata["evidence_text"].startswith(header) for node in nodes))

    def test_oversized_context_is_removed_from_search_not_metadata(self) -> None:
        title = "Very long official source title " * 8
        header = "Very long repeated record header " * 8
        document = source_document(
            f"{header}\n" + ("evidence " * 80),
            record_kind="record",
            source_title=title,
            record_header=header,
            repeat_context=header,
            structure_path=["Very long heading " * 8],
        )

        nodes = chunk_documents(
            [document], tokenizer=CharacterTokenizer(160), max_tokens=160
        )

        self.assertGreater(len(nodes), 1)
        self.assertTrue(all(node.metadata["token_count"] <= 160 for node in nodes))
        self.assertTrue(all(node.metadata["source_title"] == title for node in nodes))
        self.assertTrue(all(node.metadata["record_header"] == header for node in nodes))

    def test_search_text_whitelist(self) -> None:
        search = build_search_text({
            "source_title": "Title",
            "structure_path": ["Policy", "Borrowing"],
            "search_aliases": ["借閱"],
            "record_header": "Renewal",
            "chunk_id": "uuid-not-searchable",
            "content_hash": "hash-not-searchable",
        }, "Evidence")
        self.assertEqual(search, "Title\n\nPolicy\n\nBorrowing\n\n借閱\n\nRenewal\n\nEvidence")

    def test_unlabelled_single_uses_512_64_fallback(self) -> None:
        document = source_document("word " * 120, record_kind="auto")
        nodes = chunk_documents(
            [document], tokenizer=CharacterTokenizer(180), max_tokens=180
        )
        self.assertGreater(len(nodes), 1)
        self.assertTrue(all(node.metadata["chunk_overlap"] == 64 for node in nodes))
        self.assertTrue(all(node.metadata["chunk_policy"] == "oversized_leaf" for node in nodes))
        self.assertTrue(all(node.metadata["classification_source"] == "fallback" for node in nodes))
        for previous, following in zip(nodes, nodes[1:]):
            self.assertEqual(
                previous.metadata["evidence_text"][-64:],
                following.metadata["evidence_text"][:64],
            )


class DeterministicReaderTests(unittest.TestCase):
    def _load(self, path: Path, **kwargs: object) -> list[Document]:
        return load_file(
            path,
            document_id="fixture-doc",
            original_file_name=path.name,
            document_type=str(kwargs.pop("document_type", "auto")),
            branch_ids=kwargs.pop("branch_ids", None),
            **kwargs,
        )

    def test_csv_rows_override_faq_hint(self) -> None:
        documents = self._load(FIXTURES / "records.csv", document_type="faq")
        self.assertEqual(len(documents), 2)
        self.assertTrue(all(doc.metadata["record_kind"] == "table" for doc in documents))
        self.assertEqual(documents[0].metadata["locator"]["row"], 2)

    def test_json_records_preserve_supplied_branch_ids(self) -> None:
        documents = self._load(
            FIXTURES / "branch_records.json",
            document_type="record",
            branch_ids=["HKCL"],
        )
        self.assertEqual(len(documents), 2)
        self.assertEqual(documents[0].metadata["branch_ids"], ["HKCL"])

    def test_xml_records_have_stable_locators(self) -> None:
        documents = self._load(FIXTURES / "records.xml")
        self.assertEqual(len(documents), 2)
        self.assertNotEqual(documents[0].metadata["locator"], documents[1].metadata["locator"])

    def test_multi_sheet_xlsx_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "branches.xlsx"
            workbook = openpyxl.Workbook()
            first = workbook.active
            first.title = "Hong Kong"
            first.append(["branch", "hours"])
            first.append(["Central", "09:00"])
            second = workbook.create_sheet("Kowloon")
            second.append(["branch", "hours"])
            second.append(["Kowloon", "10:00"])
            workbook.save(path)
            documents = self._load(path)
        self.assertEqual(len(documents), 2)
        self.assertEqual(
            {document.metadata["sheet_name"] for document in documents},
            {"Hong Kong", "Kowloon"},
        )

    def test_docling_failure_is_not_silently_reparsed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.txt"
            path.write_text("structured source text", encoding="utf-8")
            with patch("src.ingestion.readers._load_docling", side_effect=RuntimeError("failed")):
                with self.assertRaisesRegex(RuntimeError, "failed"):
                    self._load(path)

    def test_empty_docling_html_uses_text_fallback_but_form_shell_stays_empty(self) -> None:
        recommended_html = """
        <div id="content">
          <div class="left-menu">
            <a href="javascript:void(0)"
               onclick="window.pdf_download='/en/common/attachments/earth.pdf'">
              About the Earth
            </a>
            <div style="display:none">
              <h2>Introduction</h2>
              <p>The Earth has mountains, oceans, deserts, and diverse wildlife.</p>
            </div>
          </div>
          <div class="inner-body"><div id="load_booklist"></div></div>
        </div>
        """
        form_shell = """
        <div id="content"><form action="/patron/login">
          <label for="password">Password</label>
          <input id="password" type="password"><button>Submit</button>
        </form></div>
        """

        with tempfile.TemporaryDirectory() as directory:
            recommended_path = Path(directory) / "recommended.html"
            recommended_path.write_text(recommended_html, encoding="utf-8")
            shell_path = Path(directory) / "shell.html"
            shell_path.write_text(form_shell, encoding="utf-8")
            with patch("src.ingestion.readers._load_docling", return_value=[]):
                recommended = self._load(
                    recommended_path,
                    source_url="https://www.hkpl.gov.hk/en/kids/recommended/reading.html",
                )
                shell = self._load(
                    shell_path,
                    source_url="https://www.hkpl.gov.hk/en/change_my_password.html",
                )

        self.assertEqual(len(recommended), 1)
        self.assertIn("About the Earth", recommended[0].text)
        self.assertIn("mountains, oceans, deserts", recommended[0].text)
        self.assertIn(
            "https://www.hkpl.gov.hk/en/common/attachments/earth.pdf",
            recommended[0].text,
        )
        self.assertEqual(shell, [])

    def test_forms_landing_page_keeps_link_records(self) -> None:
        html = """
        <table><tr><th>Form No.</th><th>Form Name</th></tr>
        <tr><td>LCS 050</td><td><a href="/forms/LCS050.pdf">Application</a></td></tr>
        </table>
        """
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "forms.html"
            path.write_text(html, encoding="utf-8")
            documents = self._load(
                path,
                document_type="prose",
                source_url="https://www.hkpl.gov.hk/en/about-us/forms.html",
            )
        self.assertEqual(len(documents), 1)
        self.assertEqual(documents[0].metadata["record_kind"], "table")
        self.assertIn("https://www.hkpl.gov.hk/forms/LCS050.pdf", documents[0].text)

    def test_event_page_keeps_field_rows_together(self) -> None:
        html = """
        <div class="main_content"><table>
          <tr><th>Date &amp; Time:</th><td>30 August 2026, 2:30 p.m.</td></tr>
          <tr><th>Venue:</th><td>Ping Shan Tin Shui Wai Public Library</td></tr>
          <tr><th>Description:</th><td>A useful public talk.</td></tr>
        </table></div>
        """
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "event.html"
            path.write_text(html, encoding="utf-8")
            documents = self._load(
                path,
                document_type="record",
                source_url="https://www.hkpl.gov.hk/en/extension-activities/event/123/example",
            )
        self.assertEqual(len(documents), 1)
        self.assertIn("Date & Time: 30 August 2026", documents[0].text)
        self.assertIn("Venue: Ping Shan", documents[0].text)

    def test_hours_table_expands_rowspans_and_repeats_header(self) -> None:
        html = """
        <div class="main_content"><table>
          <tr><th>Day</th><th>Session</th><th>Location</th></tr>
          <tr><td rowspan="2">Monday</td><td>AM</td><td>Central</td></tr>
          <tr><td>PM</td><td>Kowloon</td></tr>
        </table></div>
        """
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "hours.html"
            path.write_text(html, encoding="utf-8")
            documents = self._load(
                path,
                document_type="prose",
                source_url="https://www.hkpl.gov.hk/en/locations/opening-hours-03.html",
            )
        self.assertEqual(len(documents), 2)
        self.assertIn("Monday | PM | Kowloon", documents[1].text)
        self.assertTrue(all("Day | Session | Location" in doc.text for doc in documents))


class MainContentSelectionTests(unittest.TestCase):
    @patch("scripts.crawl_hkpl_site.save_hash")
    @patch("scripts.crawl_hkpl_site.ingest_path_sync")
    @patch("scripts.crawl_hkpl_site.classify_batch_items_resilient_sync")
    @patch("scripts.crawl_hkpl_site.extracted_classifier_text")
    def test_crawler_batch_runs_full_ingestion_handoff(
        self,
        extract_text,
        classify,
        ingest,
        save_hash,
    ) -> None:
        extract_text.return_value = "Official facts"
        classify.return_value = {
            "https://www.hkpl.gov.hk/en/page.html": {
                "document_type": "prose",
                "classification_source": "fallback",
            }
        }
        ingest.return_value = {"status": "completed"}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "page.html"
            path.write_text("<main>Official facts</main>", encoding="utf-8")
            pending = [{
                "path": path,
                "url": "https://www.hkpl.gov.hk/en/page.html",
                "title": "Official page",
                "extension": ".html",
                "mime_type": "text/html",
                "content_hash": "hash-1",
                "is_pdf": False,
                "replace_document_id": "existing-id",
            }]
            stats = {
                "failed": 0,
                "indexed": 0,
                "discovery_only": 0,
                "html_indexed": 0,
                "pdf_indexed": 0,
            }

            flush_pending(pending, stats)

        self.assertEqual(pending, [])
        self.assertEqual(stats["indexed"], 1)
        self.assertEqual(stats["html_indexed"], 1)
        self.assertEqual(
            ingest.call_args.kwargs["classification_source"],
            "fallback",
        )
        self.assertEqual(ingest.call_args.kwargs["document_type"], "prose")
        self.assertEqual(
            ingest.call_args.kwargs["replace_document_id"],
            "existing-id",
        )
        save_hash.assert_called_once_with(
            "https://www.hkpl.gov.hk/en/page.html",
            "hash-1",
        )

    def test_decodes_utf8_response_bytes_instead_of_mojibake_text(self) -> None:
        expected = "香港公共圖書館"

        class Response:
            content = f"<p>{expected}</p>".encode("utf-8")
            text = content.decode("latin-1")
            headers = {"content-type": "text/html; charset=iso-8859-1"}

        self.assertNotIn(expected, Response.text)
        self.assertIn(expected, decode_response_text(Response()))

    def test_discovers_allowed_pdf_urls_in_attributes(self) -> None:
        html = """
        <a href="/en/ordinary.html">Ordinary page</a>
        <a href="javascript:void(0)"
           onclick="download('/en/common/attachments/onclick.pdf')">PDF</a>
        <button data-download-url="/en/common/attachments/data.pdf">PDF</button>
        <button data-download-url="https://example.com/external.pdf">External</button>
        <button data-download-url="/patron/private.pdf">Private</button>
        <button data-download-url="/en/assets/app.js">Script</button>
        """

        links = discover_links(
            "https://www.hkpl.gov.hk/en/kids/recommended/reading.html",
            html,
            include_query_urls=False,
        )

        self.assertEqual(links, [
            "https://www.hkpl.gov.hk/en/common/attachments/data.pdf",
            "https://www.hkpl.gov.hk/en/common/attachments/onclick.pdf",
            "https://www.hkpl.gov.hk/en/ordinary.html",
        ])

    def test_prefers_substantive_main_content_over_large_link_sidebar(self) -> None:
        html = """
        <html><head><title>FAQ</title></head><body><div id="content">
          <div class="left_nav">{links}</div>
          <div class="main_content"><h1>Useful FAQ</h1><p>{body}</p></div>
        </div></body></html>
        """.format(
            links=" ".join(f'<a href="/{i}">Menu {i}</a>' for i in range(100)),
            body="Useful factual answer. " * 15,
        )
        _, extracted = extract_main_html(html)
        self.assertIn("Useful factual answer", extracted)
        self.assertNotIn("Menu 99", extracted)

    def test_falls_back_to_outer_content_when_inner_container_is_empty(self) -> None:
        html = """
        <html><head><title>Kids reading</title></head><body><div id="content">
          <div class="left-menu"><p>{body}</p></div><div class="inner-body"></div>
        </div></body></html>
        """.format(body="Recommended reading content. " * 12)
        _, extracted = extract_main_html(html)
        self.assertIn("Recommended reading content", extracted)


class RecordMetadataTests(unittest.TestCase):
    def test_notice_and_event_dates_are_typed_metadata(self) -> None:
        notice = extract_html_record_metadata(
            (FIXTURES / "notice_detail.html").read_text(encoding="utf-8"),
            "record",
        )
        event = extract_html_record_metadata(
            (FIXTURES / "event_detail.html").read_text(encoding="utf-8"),
            "record",
        )
        self.assertEqual(notice["datePublished"], "2026-08-01")
        self.assertEqual(event["startDate"], "2026-09-10T15:00:00+08:00")
        self.assertEqual(event["endDate"], "2026-09-10T16:00:00+08:00")


if __name__ == "__main__":
    unittest.main()
